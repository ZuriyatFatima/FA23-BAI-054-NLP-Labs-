"""Generates the sample PDF and DOCX files used by the Lab 02 tasks.
Replace the files in ./data with your own documents if you prefer."""
from pathlib import Path
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, PageBreak, Spacer

DATA = Path(__file__).parent / "data"
DATA.mkdir(exist_ok=True)

PAGES = [
    ("Introduction to Natural Language Processing",
     ["Natural language processing (NLP) is a field of artificial intelligence that focuses on the interaction between computers and human language. It is used in search engines, chatbots and translation systems. Visit https://example.com/nlp for more details.",
      "The field began in the 1950s and has grown rapidly since 2010. Today, NLP powers applications that millions of people use every day. Language models learn patterns from large text corpora."]),
    ("Text Preprocessing",
     ["Text preprocessing is the first step in any NLP pipeline. Common steps include tokenization, lowercasing, stopword removal and stemming. Clean data leads to better models.",
      "Tokenization splits text into words or sentences. Stopwords such as the, is and of are frequent but carry little meaning. Removing them reduces noise in the data."]),
    ("Feature Extraction",
     ["Machine learning models cannot read raw text, so text must be converted into numbers. Bag of words and TF-IDF are classic feature extraction methods. Word embeddings capture semantic meaning in dense vectors.",
      "Modern language models use contextual embeddings. These representations change depending on the surrounding words, which helps with ambiguity in language."]),
    ("Machine Learning for Text",
     ["Classification, clustering and sequence labelling are common tasks. Sentiment analysis predicts whether a review is positive or negative. Named entity recognition finds people, places and organizations in text.",
      "Neural networks, especially transformers, have achieved state of the art results. Training these models requires large datasets and powerful hardware."]),
    ("Evaluation of NLP Systems",
     ["Evaluation metrics include accuracy, precision, recall and the F1 score. For translation, BLEU is widely used. Human evaluation is still important for tasks like summarization and dialogue.",
      "A good evaluation uses a held-out test set. Cross validation can give a more reliable estimate of performance on small datasets."]),
    ("Applications and Future Directions",
     ["NLP applications include virtual assistants, spam filters, machine translation and document summarization. Language technology is improving access to information for people worldwide.",
      "Future directions include multilingual models, low-resource languages and more efficient training. Ethical issues such as bias and fairness must also be addressed in language technology."]),
]

def make_pdf():
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(DATA / "sample_report.pdf"), pagesize=A4)
    story = []
    for i, (title, paras) in enumerate(PAGES):
        story.append(Paragraph(title, styles["Heading1"]))
        for p in paras:
            story += [Paragraph(p, styles["BodyText"]), Spacer(1, 12)]
        if i < len(PAGES) - 1:
            story.append(PageBreak())
    doc.build(story)

def make_docx_a():
    d = Document()
    d.add_heading("NLP Overview", level=0)
    d.add_paragraph("This document introduces the core ideas of natural language processing and its applications in industry.")
    d.add_heading("Introduction", level=1)
    d.add_paragraph("Natural language processing enables computers to understand human language. It combines linguistics, computer science and machine learning.")
    d.add_paragraph("Language data is messy, so careful preprocessing is required before any model is trained.")
    d.add_paragraph("")
    d.add_heading("Tokenization", level=1)
    d.add_paragraph("Tokenization splits text into smaller units called tokens. Tokens can be words, subwords or characters depending on the model.")
    d.add_paragraph("Short note.")
    d.add_heading("Results", level=1)
    d.add_paragraph("The table below summarises the accuracy of three models on the test set.")
    t = d.add_table(rows=4, cols=3)
    t.style = "Table Grid"
    rows = [("Model", "Accuracy", "F1 Score"), ("Naive Bayes", "0.82", "0.80"),
            ("SVM", "0.88", "0.87"), ("Transformer", "0.94", "")]
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            t.cell(r, c).text = val
    d.add_paragraph("Transformers achieved the best performance, although they need more computing resources.")
    d.add_heading("Conclusion", level=1)
    d.add_paragraph("Language data is messy, so careful preprocessing is required before any model is trained.")
    d.save(DATA / "doc_a.docx")

def make_docx_b():
    d = Document()
    d.add_heading("Text Mining Notes", level=0)
    d.add_paragraph("These notes cover practical text mining techniques for beginners in data science.")
    d.add_heading("Introduction", level=1)
    d.add_paragraph("Natural language processing enables computers to understand human language. It combines linguistics, computer science and machine learning.")
    d.add_paragraph("Word frequency analysis is a simple yet powerful way to explore a corpus and find its most important terms.")
    d.add_heading("Stopword Removal", level=1)
    d.add_paragraph("Stopwords are common words that add little meaning, and removing them helps models focus on informative terms.")
    d.add_heading("Conclusion", level=1)
    d.add_paragraph("Practical text mining combines careful cleaning, sensible features and honest evaluation of the final results.")
    d.save(DATA / "doc_b.docx")

if __name__ == "__main__":
    make_pdf(); make_docx_a(); make_docx_b()
    print("Sample files created in", DATA)
