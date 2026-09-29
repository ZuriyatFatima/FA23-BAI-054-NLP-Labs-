"""Lab 02 - Task 6: Multi-Document Comparison, Deduplication, and Frequency Report

Usage: python task_06_docx_comparison.py [docx_A] [docx_B]
"""
import re
import sys
from collections import Counter
from pathlib import Path

from docx import Document

DATA = Path(__file__).parent / "data"
PATH_A = sys.argv[1] if len(sys.argv) > 1 else DATA / "doc_a.docx"
PATH_B = sys.argv[2] if len(sys.argv) > 2 else DATA / "doc_b.docx"
doc_a, doc_b = Document(PATH_A), Document(PATH_B)

STOPWORDS = {"the", "is", "a", "an", "and", "or", "of", "to", "in", "on", "for", "it", "as",
             "are", "was", "be", "by", "with", "that", "this", "at", "so", "them", "they"}


def headings(doc):
    return [p.text.strip() for p in doc.paragraphs
            if p.style.name.startswith("Heading") and p.text.strip()]


def normal_paragraphs(doc):
    # blank paragraphs are ignored so they don't count as "duplicates"
    return [p.text for p in doc.paragraphs if p.style.name == "Normal" and p.text.strip()]


# (a) heading comparison using set operations
head_a, head_b = headings(doc_a), headings(doc_b)
set_a, set_b = set(head_a), set(head_b)
print("(a-i)   Common headings   :", sorted(set_a & set_b))
print("(a-ii)  Only in Document A:", sorted(set_a - set_b))
print("(a-iii) Only in Document B:", sorted(set_b - set_a))

# (b) merge + deduplicate Normal paragraphs
merged = normal_paragraphs(doc_a) + normal_paragraphs(doc_b)
seen, unique = set(), []
for p in merged:
    key = p.strip().lower()
    if key not in seen:
        seen.add(key)
        unique.append(p)
print(f"\n(b) Paragraphs before deduplication: {len(merged)} | after: {len(unique)}")

# (c) word frequencies
freq = Counter()
tokenised = []
for p in unique:
    words = re.sub(r"[^\w\s]|\d|_", "", p.lower()).split()
    tokenised.append(words)
    freq.update(w for w in words if w not in STOPWORDS)
print(f"\n(c) Top 20 words (after removing {len(STOPWORDS)} stopwords):")
for word, count in freq.most_common(20):
    print(f"    {word:15} {count}")

# (d) summary report
lengths = [len(w) for w in tokenised]
longest = max(unique, key=lambda p: len(p.split()))
vocab_size = len({w for words in tokenised for w in words})
print("\n(d) SUMMARY")
print("    (i)   Total headings (both docs)   :", len(head_a) + len(head_b))
print("    (ii)  Unique body paragraphs       :", len(unique))
print("    (iii) Vocabulary size (word types) :", vocab_size)
print("    (iv)  Average paragraph length     :", round(sum(lengths) / len(lengths), 2), "words")
print(f"    (v)   Longest paragraph ({len(longest.split())} words):\n          {longest}")
